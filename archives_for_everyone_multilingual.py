"""
=============================================================
  Archives-for-Everyone | OCR → DOCX PIPELINE (2 columns)
  Author: Alessandro Ceppetelli
  Configurable source languages → configurable target language
=============================================================

INSTALLATION (run once in the terminal):
  pip install pytesseract pillow argostranslate python-docx opencv-python langdetect pillow-heif pdf2image pdfplumber

Tesseract (OCR engine) — download and install from:
  https://github.com/UB-Mannheim/tesseract/wiki
  Install the Tesseract language packs matching your selected source languages

EXPECTED FOLDER STRUCTURE:
  - One main folder
  - Containing folders and subfolders
  - Each subfolder containing images/PDFs generates one .docx file
  - Usage: python archives_for_everyone.py --source nl,fr,de --target it --input ./input --output ./output
  - The .docx filename matches the subfolder name
=============================================================
"""

import os
import sys
import csv
import glob
import textwrap
import argparse
import shutil
from collections import deque

# ── dependencies ──────────────────────────────────────────────
try:
    import cv2
    import numpy as np
    import pytesseract
    from PIL import Image
    from pillow_heif import register_heif_opener
    register_heif_opener()
    from pdf2image import convert_from_path
    from langdetect import detect
    import argostranslate.package
    import argostranslate.translate
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError as e:
    print(f"\n❌ Missing module: {e}")
    print("Run: pip install pytesseract pillow argostranslate python-docx opencv-python langdetect pillow-heif pdf2image pdfplumber")
    sys.exit(1)


# =============================================================
#  CONFIGURATION
# =============================================================

ARCHIVE_DIR = os.path.join(os.getcwd(), "input")
OUTPUT_DIR = os.path.join(os.getcwd(), "output")
IMG_EXT = ("heic", "jpg", "jpeg", "png", "pdf")

# ISO 639-1 language codes mapped to Tesseract traineddata identifiers.
# More languages can be added if the corresponding Tesseract packs are installed.
TESSERACT_CODES = {
    "ar": "ara", "bg": "bul", "ca": "cat", "cs": "ces", "da": "dan",
    "de": "deu", "el": "ell", "en": "eng", "es": "spa", "et": "est",
    "fi": "fin", "fr": "fra", "he": "heb", "hi": "hin", "hr": "hrv",
    "hu": "hun", "id": "ind", "it": "ita", "ja": "jpn", "ko": "kor",
    "lt": "lit", "lv": "lav", "nl": "nld", "no": "nor", "pl": "pol",
    "pt": "por", "ro": "ron", "ru": "rus", "sk": "slk", "sl": "slv",
    "sv": "swe", "th": "tha", "tr": "tur", "uk": "ukr", "vi": "vie",
    "zh": "chi_sim",
}
DEFAULT_SOURCES = ["nl", "fr", "en", "de"]
SOURCE_LANGUAGES = DEFAULT_SOURCES.copy()
TARGET_LANGUAGE = "en"
TRANSLATION_ROUTES = {}


# ── theme colors ──────────────────────────────────────────────
COLOR_HEADER_BG  = "2E5FA3"
COLOR_ORIG_BG    = "EBF2FB"
COLOR_TRANS_BG   = "F0F7EE"
COLOR_HEADER_TXT = "FFFFFF"


def find_translation_route(source, target, edges):
    """Find the shortest available Argos package chain (may use pivot languages)."""
    if source == target:
        return []
    queue = deque([(source, [])])
    visited = {source}
    while queue:
        current, route = queue.popleft()
        for nxt in sorted(edges.get(current, [])):
            if nxt in visited:
                continue
            new_route = route + [(current, nxt)]
            if nxt == target:
                return new_route
            visited.add(nxt)
            queue.append((nxt, new_route))
    return None


def install_argos_models():
    """Install packages for each chosen source→target route, when available."""
    global TRANSLATION_ROUTES
    sources = [code for code in SOURCE_LANGUAGES if code != TARGET_LANGUAGE]
    if not sources:
        return
    print("🔄 Checking available Argos Translate language pairs...")
    argostranslate.package.update_package_index()
    packages = argostranslate.package.get_available_packages()
    available_by_pair = {(p.from_code, p.to_code): p for p in packages}
    # Read direct installed model packages, not derived/pivot translations.
    # Argos Language objects have translations_from/translations_to, not translations.
    installed_pairs = {
        (pkg.from_code, pkg.to_code)
        for pkg in argostranslate.package.get_installed_packages()
        if getattr(pkg, "from_code", None) and getattr(pkg, "to_code", None)
        and getattr(pkg, "type", "translate") == "translate"
    }
    edges = {}
    for src, dst in set(available_by_pair) | installed_pairs:
        edges.setdefault(src, set()).add(dst)
    for source in sources:
        route = find_translation_route(source, TARGET_LANGUAGE, edges)
        if route is None:
            print(f"⚠️  No Argos package route for {source} → {TARGET_LANGUAGE}. Translation will be unavailable.")
            continue
        TRANSLATION_ROUTES[source] = route
        for pair in route:
            if pair not in installed_pairs:
                pkg = available_by_pair.get(pair)
                if pkg is None:
                    print(f"⚠️  Missing package {pair[0]} → {pair[1]}")
                    continue
                print(f"⬇️  Installing {pair[0]} → {pair[1]}...")
                argostranslate.package.install_from_path(pkg.download())
                installed_pairs.add(pair)
    print("✅ Translation package check completed.")


def configure_tesseract(explicit_path=None):
    """Locate Tesseract from a CLI option, environment, PATH or common Windows folders."""
    candidates = []
    if explicit_path:
        candidates.append(explicit_path)
    environment_path = os.environ.get("TESSERACT_CMD")
    if environment_path:
        candidates.append(environment_path)
    on_path = shutil.which("tesseract")
    if on_path:
        candidates.append(on_path)
    if os.name == "nt":
        home = os.path.expanduser("~")
        for base in (os.environ.get("LOCALAPPDATA"),
                     os.environ.get("ProgramFiles"),
                     os.environ.get("ProgramFiles(x86)"),
                     os.path.join(home, "AppData", "Local", "Programs")):
            if base:
                candidates.append(os.path.join(base, "Tesseract-OCR", "tesseract.exe"))
    for candidate in candidates:
        candidate = os.path.expandvars(os.path.expanduser(candidate.strip('"')))
        resolved = shutil.which(candidate) or candidate
        if os.path.isfile(resolved):
            pytesseract.pytesseract.tesseract_cmd = resolved
            print(f"✅ Tesseract found: {resolved}")
            return resolved
    raise RuntimeError(
        "Tesseract OCR executable not found. Install it from "
        "https://github.com/UB-Mannheim/tesseract/wiki and ensure it is on PATH, "
        "or pass --tesseract \"C:\\path\\to\\tesseract.exe\". "
        "You can also set the TESSERACT_CMD environment variable."
    )


def validate_ocr_languages():
    """Fail early if a selected Tesseract OCR model is missing."""
    installed = set(pytesseract.get_languages(config=""))
    missing = [f"{code} ({TESSERACT_CODES[code]})" for code in SOURCE_LANGUAGES
               if TESSERACT_CODES[code] not in installed]
    if missing:
        raise RuntimeError("Missing Tesseract language packs: " + ", ".join(missing))


def select_source_language(text, previous_language=None):
    """Keep detection within the selected source-language set."""
    if len(SOURCE_LANGUAGES) == 1:
        return SOURCE_LANGUAGES[0]
    if not text or len(text.strip()) < 20:
        return previous_language if previous_language in SOURCE_LANGUAGES else SOURCE_LANGUAGES[0]
    try:
        from langdetect import detect_langs
        candidates = detect_langs(text)
        allowed = [candidate for candidate in candidates if candidate.lang in SOURCE_LANGUAGES]
        if allowed:
            return max(allowed, key=lambda item: item.prob).lang
    except Exception:
        pass
    return previous_language if previous_language in SOURCE_LANGUAGES else SOURCE_LANGUAGES[0]


def configure_arguments(argv=None):
    parser = argparse.ArgumentParser(description="Archives-for-Everyone: OCR and multilingual archival translation")
    parser.add_argument("--source", default="auto", help="Source codes separated by commas (e.g. nl,fr,it), or auto (NL,FR,EN,DE)")
    parser.add_argument("--target", default="en", help="Translation target (ISO language code, e.g. en,fr,it,es)")
    parser.add_argument("--input", default=ARCHIVE_DIR, help="Input archive root directory")
    parser.add_argument("--output", default=OUTPUT_DIR, help="DOCX output directory")
    parser.add_argument("--tesseract", default=None, help="Optional path to the tesseract executable")
    parser.add_argument("--interactive", action="store_true", help="Ask for source and target languages interactively")
    parser.add_argument("--list-languages", action="store_true", help="List recognized source-language codes")
    args = parser.parse_args(argv)
    if args.list_languages:
        print("Supported source codes (requires corresponding Tesseract traineddata):")
        print(", ".join(sorted(TESSERACT_CODES)))
        return None
    if args.interactive:
        args.source = input("Source languages [auto or comma-separated codes, default auto]: ").strip() or "auto"
        args.target = input("Target language [code, default en]: ").strip() or "en"
    args.target = args.target.strip().lower()
    args.source_codes = (DEFAULT_SOURCES.copy() if args.source.strip().lower() == "auto"
                         else list(dict.fromkeys(x.strip().lower() for x in args.source.split(",") if x.strip())))
    unknown = [code for code in args.source_codes if code not in TESSERACT_CODES]
    if not args.source_codes or unknown:
        parser.error("Unknown/empty source languages: " + ", ".join(unknown))
    if not args.target or len(args.target) < 2:
        parser.error("Specify a valid target language code")
    return args


def preprocess_image(image_path):
    """Improves image quality for OCR."""
    pil_img = Image.open(image_path).convert("RGB")
    img = cv2.cvtColor(np.array(pil_img), cv2.COLOR_RGB2BGR)

    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)

    coords = np.column_stack(np.where(gray < 200))
    if len(coords) > 100:
        angle = cv2.minAreaRect(coords)[-1]
        if angle < -45:
            angle = -(90 + angle)
        else:
            angle = -angle
        if abs(angle) > 0.5:
            (h, w) = gray.shape
            M = cv2.getRotationMatrix2D((w // 2, h // 2), angle, 1.0)
            gray = cv2.warpAffine(gray, M, (w, h), flags=cv2.INTER_CUBIC,
                                  borderMode=cv2.BORDER_REPLICATE)

    gray = cv2.GaussianBlur(gray, (3, 3), 0)
    _, processed = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    processed = cv2.medianBlur(processed, 3)
    processed = cv2.morphologyEx(processed, cv2.MORPH_OPEN, np.ones((2, 2), np.uint8))

    return Image.fromarray(processed)


def ocr_image(image_path, previous_language=None):
    """
    Basic OCR on a single image (without column handling).
    """
    img = preprocess_image(image_path)

    language_order = [TESSERACT_CODES[code] for code in SOURCE_LANGUAGES]
    if previous_language in SOURCE_LANGUAGES:
        preferred = TESSERACT_CODES[previous_language]
        language_order = [preferred] + [code for code in language_order if code != preferred]

    text_candidates = []

    for lang in language_order:
        try:
            config = "--psm 4" if lang == "nld" else "--psm 3"
            text = pytesseract.image_to_string(img, lang=lang, config=config).strip()
            if text:
                text_candidates.append(text)
        except Exception:
            continue

    if not text_candidates:
        return ""

    def score(t):
        letters = sum(c.isalpha() for c in t)
        words = len(t.split())
        return letters + words * 2

    text_candidates.sort(key=score, reverse=True)
    return text_candidates[0]


def ocr_with_layout(image_path, previous_language=None):
    """OCR using a selected language and Tesseract text-block layout detection."""
    img = preprocess_image(image_path)
    # First-page documents use the selected language models jointly for OCR.
    # Following pages reuse the previously detected language for improved accuracy.
    if previous_language in SOURCE_LANGUAGES:
        lang = TESSERACT_CODES[previous_language]
    else:
        lang = "+".join(TESSERACT_CODES[code] for code in SOURCE_LANGUAGES)

    try:
        data = pytesseract.image_to_data(
            img,
            lang=lang,
            config="--psm 1",
            output_type=pytesseract.Output.DICT
        )
    except Exception:
        return ocr_image(image_path, previous_language)

    blocks = {}
    n = len(data["text"])

    for i in range(n):
        txt = (data["text"][i] or "").strip()
        if not txt:
            continue

        try:
            conf = float(data["conf"][i])
        except Exception:
            conf = -1

        if conf < 8:
            continue

        if sum(c.isalpha() for c in txt) == 0:
            continue

        block_num = data["block_num"][i]
        x = int(data["left"][i])
        y = int(data["top"][i])
        w = int(data["width"][i])
        h = int(data["height"][i])

        if block_num not in blocks:
            blocks[block_num] = {
                "x1": x,
                "y1": y,
                "x2": x + w,
                "y2": y + h,
                "words": []
            }
        else:
            blocks[block_num]["x1"] = min(blocks[block_num]["x1"], x)
            blocks[block_num]["y1"] = min(blocks[block_num]["y1"], y)
            blocks[block_num]["x2"] = max(blocks[block_num]["x2"], x + w)
            blocks[block_num]["y2"] = max(blocks[block_num]["y2"], y + h)

        blocks[block_num]["words"].append({
            "text": txt,
            "line_num": data["line_num"][i],
            "par_num": data["par_num"][i],
            "x": x,
            "y": y
        })

    if not blocks:
        return ocr_image(image_path, previous_language)

    image_width, image_height = img.size
    valid_blocks = []

    for _, b in blocks.items():
        width = b["x2"] - b["x1"]
        height = b["y2"] - b["y1"]
        word_count = len(b["words"])
        letter_count = sum(sum(c.isalpha() for c in p["text"]) for p in b["words"])

        if word_count < 5:
            continue
        if letter_count < 20:
            continue
        if height < int(image_height * 0.05):
            continue
        if width > int(image_width * 0.75):
            continue

        valid_blocks.append(b)

    if not valid_blocks:
        return ocr_image(image_path, previous_language)

    valid_blocks.sort(key=lambda b: b["x1"])

    columns = []
    column_threshold = max(40, int(image_width * 0.08))

    for b in valid_blocks:
        if not columns:
            columns.append([b])
            continue

        mean_x = sum(x["x1"] for x in columns[-1]) / len(columns[-1])
        if abs(b["x1"] - mean_x) <= column_threshold:
            columns[-1].append(b)
        else:
            columns.append([b])

    results = []

    for column in columns:
        column.sort(key=lambda b: b["y1"])

        for b in column:
            lines = {}
            for p in b["words"]:
                key = (p["par_num"], p["line_num"])
                if key not in lines:
                    lines[key] = []
                lines[key].append(p)

            sorted_lines = []
            for key, line_words in lines.items():
                line_words.sort(key=lambda x: x["x"])
                sorted_lines.append(
                    (min(p["y"] for p in line_words),
                     " ".join(p["text"] for p in line_words))
                )

            sorted_lines.sort(key=lambda x: x[0])

            block_text = "\n".join(r[1] for r in sorted_lines).strip()
            block_text = clean_xml_text(join_hyphenated_words(block_text))

            if sum(c.isalpha() for c in block_text) >= 20:
                results.append(block_text)

    if results:
        return "\n\n".join(results)

    return ocr_image(image_path, previous_language)


def join_hyphenated_words(text):
    """Joins words split by line-ending hyphens."""
    return text.replace("-\n", "")

def prepare_text_for_translation(text):
    """
    Combines OCR lines into paragraphs before translation.
    Ends a paragraph only when strong punctuation is found.
    """
    if not text:
        return ""

    text = join_hyphenated_words(text)
    lines = [r.strip() for r in text.splitlines()]

    paragraphs = []
    current_paragraph = []

    for line in lines:
        if not line:
            if current_paragraph:
                paragraphs.append(" ".join(current_paragraph))
                current_paragraph = []
            continue

        current_paragraph.append(line)

        if line.endswith((".", "!", "?", ":", ";")):
            paragraphs.append(" ".join(current_paragraph))
            current_paragraph = []

    if current_paragraph:
        paragraphs.append(" ".join(current_paragraph))

    return "\n\n".join(paragraphs)

def clean_xml_text(text):
    """Removes characters incompatible with XML/Word."""
    if not text:
        return ""
    return "".join(
        c for c in text
        if c == "\n" or c == "\t" or ord(c) >= 32
    )


def is_real_table(table):
    """
    Determines whether a table extracted by pdfplumber is a genuine table.
    Avoids false positives: ordinary text blocks misidentified as single-column tables.
    """
    if not table:
        return False

    lines = []
    for row in table:
        if not row:
            continue
        cells = [(c or "").strip() for c in row]
        if any(cells):
            lines.append(cells)

    if len(lines) < 2:
        return False

    max_cols = max(len(r) for r in lines)

    # If there is only one column, it is almost never a real table
    if max_cols < 2:
        return False

    # Count rows with at least two nonempty cells
    rows_with_two_cells = 0
    for r in lines:
        nonempty_cells = [c for c in r if c.strip()]
        if len(nonempty_cells) >= 2:
            rows_with_two_cells += 1

    if rows_with_two_cells < 2:
        return False

    return True


def extract_pdf_text_and_tables(pdf_path):
    """
    Extracts content from text-based PDFs.
    - Real tables are returned as structured tables.
    - False positives from pdfplumber are ignored.
    - Ordinary text is extracted using the native PDF text flow.
    """
    import pdfplumber

    output_pages = []

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            page_elements = []

            # 1. Real tables
            try:
                tables = page.extract_tables()
            except Exception:
                tables = []

            for table in tables or []:
                if not is_real_table(table):
                    continue

                clean_rows = []
                for row in table:
                    if not row:
                        continue

                    cells = [(c or "").strip().replace("\n", " ") for c in row]

                    if any(cells):
                        clean_rows.append(cells)

                if clean_rows:
                    page_elements.append({
                        "type": "table",
                        "content": clean_rows
                    })

            # 2. Normal page text
            try:
                text = page.extract_text(
                    x_tolerance=2,
                    y_tolerance=3,
                    use_text_flow=True
                )
            except TypeError:
                words = page.extract_words(
                    use_text_flow=True,
                    x_tolerance=2,
                    y_tolerance=3,
                    keep_blank_chars=False
                )

                lines = []
                line = []
                last_top = None

                for w in words:
                    top = w["top"]

                    if last_top is None or abs(top - last_top) < 4:
                        line.append(w)
                        last_top = top
                    else:
                        line.sort(key=lambda x: x["x0"])
                        lines.append(" ".join(x["text"] for x in line))
                        line = [w]
                        last_top = top

                if line:
                    line.sort(key=lambda x: x["x0"])
                    lines.append(" ".join(x["text"] for x in line))

                text = "\n".join(lines)

            text = text or ""
            text = clean_xml_text(join_hyphenated_words(text)).strip()

            if text:
                page_elements.append({
                    "type": "text",
                    "content": text
                })

            # If nothing was found, add an empty page
            if not page_elements:
                page_elements.append({
                    "type": "text",
                    "content": ""
                })

            output_pages.extend(page_elements)

    return output_pages


def translate_table(table, source_language):
    """
    Translates a table cell by cell.
    """
    translated_table = []
    for row in table:
        new_row = []
        for cell in row:
            cell = clean_xml_text(join_hyphenated_words(cell or ""))
            if not cell.strip():
                new_row.append("")
            else:
                new_row.append(translate_text(cell, source_language))
        translated_table.append(new_row)
    return translated_table


def extract_pages_from_pdf(pdf_path, previous_language=None):
    """
    1) Try native PDF extraction (text and tables)
    2) Fall back to OCR only when necessary
    """
    output_pages = []

    try:
        elementi = extract_pdf_text_and_tables(pdf_path)

        if elementi:
            for i, element in enumerate(elementi, start=1):
                if element["type"] == "table":
                    table_text = "\n".join(" ".join(r) for r in element["content"])
                    table_text = clean_xml_text(join_hyphenated_words(table_text))
                    language = detect_language(table_text, previous_language)
                    previous_language = language

                    output_pages.append({
                        "number": f"{os.path.basename(pdf_path)}_p{i}",
                        "type": "table",
                        "original_table": element["content"],
                        "translated_table": translate_table(element["content"], language),
                        "language": language,
                    })
                    continue

                original_text = clean_xml_text(join_hyphenated_words(element["content"]))

                if not original_text.strip():
                    continue

                language = detect_language(original_text, previous_language)
                previous_language = language
                translated_text = translate_text(original_text, language)

                output_pages.append({
                    "number": f"{os.path.basename(pdf_path)}_p{i}",
                    "type": "text",
                    "original_text": original_text,
                    "language": language,
                    "translated_text": translated_text,
                })

            if output_pages:
                return output_pages, previous_language

    except Exception:
        pass

    try:
        pdf_pages = convert_from_path(pdf_path)
    except Exception as e:
        print(f"PDF ERROR: {e}")
        return [], previous_language

    for i, pdf_page in enumerate(pdf_pages):
        temp = f"_tmp_pdf_{i}.png"
        pdf_page.save(temp)

        try:
            original_text = clean_xml_text(join_hyphenated_words(ocr_with_layout(temp, previous_language)))
        finally:
            if os.path.exists(temp):
                os.remove(temp)

        if not original_text.strip():
            continue

        language = detect_language(original_text, previous_language)
        previous_language = language
        translated_text = translate_text(original_text, language)

        output_pages.append({
            "number": f"{os.path.basename(pdf_path)}_p{i+1}",
            "type": "text",
            "original_text": original_text,
            "language": language,
            "translated_text": translated_text,
        })

    return output_pages, previous_language


def detect_language(text, previous_language=None):
    """Detect within the user-selected languages; honour single-language mode."""
    return select_source_language(text, previous_language)


def translate_text(text, source_language):
    """Translate to the chosen target via installed Argos packages (possibly through pivots)."""
    if not text.strip() or source_language == TARGET_LANGUAGE:
        return text
    route = TRANSLATION_ROUTES.get(source_language)
    if route is None:
        return f"[No translation route available: {source_language} → {TARGET_LANGUAGE}]"
    try:
        translation_input = prepare_text_for_translation(text)
        paragraphs = translation_input.split("\n\n")
        translated = []
        for paragraph in paragraphs:
            current_text = paragraph
            if paragraph.strip():
                for src, dst in route:
                    installed = argostranslate.translate.get_installed_languages()
                    source = next(l for l in installed if l.code == src)
                    destination = next(l for l in installed if l.code == dst)
                    current_text = source.get_translation(destination).translate(current_text)
            translated.append(current_text)
        return "\n\n".join(translated)
    except Exception as e:
        return f"[Translation error ({source_language} → {TARGET_LANGUAGE}): {e}]"


def set_cell_shading(cell, hex_color):
    """Sets the background color of a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_cell_border(cell):
    """Adds thin borders to a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "BBBBBB")
        tcBorders.append(b)
    tcPr.append(tcBorders)


def create_nested_table(cell, data, font_size=8):
    """
    Inserts a real Word table inside a cell.
    """
    if not data:
        return

    rows = len(data)
    cols = max(len(r) for r in data) if data else 1

    tab = cell.add_table(rows=rows, cols=cols)
    tab.style = "Table Grid"

    for i, row in enumerate(data):
        for j in range(cols):
            text = row[j] if j < len(row) else ""
            p = tab.cell(i, j).paragraphs[0]
            run = p.add_run(clean_xml_text(str(text)))
            run.font.size = Pt(font_size)


def create_docx(document_name, pages):
    """
    Creates a .docx file with a two-column table for each page.
    Supports text pages and table pages.
    """
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)

    title = doc.add_heading(document_name, level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.color.rgb = RGBColor(0x2E, 0x5F, 0xA3)

    doc.add_paragraph("")

    for pg in pages:
        page_number = pg["number"]
        language = pg["language"].upper()

        label = language

        page_label = doc.add_paragraph()
        page_label.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = page_label.add_run(f"▌ Page {page_number}")
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x2E, 0x5F, 0xA3)

        table = doc.add_table(rows=2, cols=2)
        table.style = "Table Grid"

        for line in table.rows:
            line.cells[0].width = Cm(13)
            line.cells[1].width = Cm(13)

        header_row = table.rows[0]

        original_header_cell = header_row.cells[0]
        original_header_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(original_header_cell, COLOR_HEADER_BG)
        add_cell_border(original_header_cell)
        p = original_header_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"ORIGINAL  ({label})")
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        translation_header_cell = header_row.cells[1]
        translation_header_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(translation_header_cell, COLOR_HEADER_BG)
        add_cell_border(translation_header_cell)
        p = translation_header_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"TRANSLATION  ({TARGET_LANGUAGE.upper()})")
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        content_row = table.rows[1]

        original_cell = content_row.cells[0]
        set_cell_shading(original_cell, COLOR_ORIG_BG)
        add_cell_border(original_cell)

        translation_cell = content_row.cells[1]
        set_cell_shading(translation_cell, COLOR_TRANS_BG)
        add_cell_border(translation_cell)

        if pg.get("type") == "table":
            original_cell.paragraphs[0].clear()
            translation_cell.paragraphs[0].clear()
            create_nested_table(original_cell, pg["original_table"], font_size=8)
            create_nested_table(translation_cell, pg["translated_table"], font_size=8)
        else:
            original_text = clean_xml_text(pg["original_text"])
            translated_text = clean_xml_text(pg["translated_text"])

            original_cell.paragraphs[0].clear()
            for i, text_line in enumerate(original_text.split("\n")):
                p = original_cell.paragraphs[0] if i == 0 else original_cell.add_paragraph()
                run = p.add_run(text_line)
                run.font.size = Pt(9)
                run.font.name = "Courier New"

            translation_cell.paragraphs[0].clear()
            for i, text_line in enumerate(translated_text.split("\n")):
                p = translation_cell.paragraphs[0] if i == 0 else translation_cell.add_paragraph()
                run = p.add_run(text_line)
                run.font.size = Pt(9)

        doc.add_paragraph("")

    return doc


def number_from_filename(path):
    """Extracts the sequence number from a filename."""
    base = os.path.splitext(os.path.basename(path))[0]
    digits = "".join(filter(str.isdigit, base))
    return int(digits) if digits else 0


def find_subfolders_with_files(root_folder, ext):
    """
    Returns all subfolders containing at least one supported file.
    """
    subfolders = []

    output_root = os.path.abspath(OUTPUT_DIR)
    for root, dirs, filenames in os.walk(root_folder):
        dirs[:] = [d for d in dirs if os.path.abspath(os.path.join(root, d)) != output_root]
        matching_files = []
        for filename in filenames:
            if filename.lower().endswith(tuple(f".{e}" for e in ext)):
                matching_files.append(os.path.join(root, filename))

        if matching_files:
            matching_files.sort(key=number_from_filename)
            subfolders.append((root, matching_files))

    return subfolders


def main(argv=None):
    global ARCHIVE_DIR, OUTPUT_DIR, SOURCE_LANGUAGES, TARGET_LANGUAGE
    args = configure_arguments(argv)
    if args is None:
        return
    ARCHIVE_DIR = os.path.abspath(args.input)
    OUTPUT_DIR = os.path.abspath(args.output)
    SOURCE_LANGUAGES = args.source_codes
    TARGET_LANGUAGE = args.target
    print(f"Source languages: {', '.join(SOURCE_LANGUAGES)} | Target: {TARGET_LANGUAGE}")
    try:
        configure_tesseract(args.tesseract)
        validate_ocr_languages()
    except Exception as exc:
        print(f"❌ OCR configuration error: {exc}")
        sys.exit(2)
    print("\n" + "=" * 60)
    print("  OCR → TRANSLATION → DOCX PIPELINE")
    print("=" * 60)

    if not os.path.isdir(ARCHIVE_DIR):
        print(f"❌ Archive folder not found: {ARCHIVE_DIR}")
        sys.exit(1)

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    install_argos_models()

    groups = find_subfolders_with_files(ARCHIVE_DIR, IMG_EXT)
    print(f"\n📂 Found {len(groups)} subfolders containing images.\n")

    for folder_path, files in groups:
        document_name = os.path.basename(folder_path)
        print(f"{'─' * 50}")
        print(f"📄 Document: {document_name}")
        pages = []
        previous_language = None

        for image_path in files:
            filename = os.path.basename(image_path)

            if image_path.lower().endswith(".pdf"):
                print(f"  📄 PDF {filename}...", end=" ", flush=True)
                try:
                    pdf_pages, previous_language = extract_pages_from_pdf(image_path, previous_language)
                except Exception as e:
                    print(f"PDF ERROR: {e}")
                    continue

                if not pdf_pages:
                    print("no usable text — skipping.")
                    continue

                pages.extend(pdf_pages)
                print("extracted ✅")
                continue

            print(f"  🔍 OCR {filename}...", end=" ", flush=True)
            try:
                original_text = clean_xml_text(join_hyphenated_words(ocr_with_layout(image_path, previous_language)))
            except Exception as e:
                print(f"OCR ERROR: {e}")
                continue

            if not original_text.strip():
                print("empty text — skipping.")
                continue

            language = detect_language(original_text, previous_language)
            previous_language = language
            print(f"detected language: {language.upper()}", end=" | ", flush=True)

            translated_text = translate_text(original_text, language)

            pages.append({
                "number": filename,
                "type": "text",
                "original_text": original_text,
                "language": language,
                "translated_text": translated_text,
            })
            print("translated ✅")

        if not pages:
            print(f"  ⚠️  No pages processed for {document_name}.")
            continue

        print(f"  💾 Creating {document_name}.docx...", end=" ", flush=True)
        doc = create_docx(document_name, pages)
        output_path = os.path.join(OUTPUT_DIR, f"{document_name}.docx")
        doc.save(output_path)
        print(f"saved → {output_path} ✅")

    print(f"\n{'=' * 60}")
    print(f"✅ Completed! Files saved in: {OUTPUT_DIR}")
    print(f"{'=' * 60}\n")


if __name__ == "__main__":
    main()